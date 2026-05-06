import textwrap
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

REPORTS_DIR = Path(__file__).parent / "data" / "reports"

_QUALITY_COLOURS = {
    "missing": "FFD7D7",   # red tint
    "short":   "FFF3CD",   # amber tint
    "good":    "D8F0D8",   # green tint
    "long":    "D6EAF8",   # blue tint
}


# ── Text / email report ────────────────────────────────────────────────────────

def export_text(report: dict) -> Path:
    date_str = report["date"]
    v = report["volume"]
    d = report["descriptions"]
    p = report["pricing"]
    occ = report["occasion_products"]

    dept_lines = "\n".join(
        f"  {dept:<40} {data['url_count']:>5,}  mean £{str(data['price_mean'] or '-'):>7}  median £{str(data['price_median'] or '-'):>7}"
        for dept, data in list(report["departments"].items())[:15]
    )

    dup_lines = "\n".join(
        f"  [{ex['family_id_count']} families] {ex['pre_text'][:90] or '(no pre-content)'}"
        for ex in d["cross_parent_duplicate_examples"][:5]
    ) or "  None"

    family_issue_lines = "\n".join(
        f"  {issue['family_id']}  ({issue['distinct_description_count']} distinct descriptions)"
        for issue in d["family_description_issues"][:10]
    ) or "  None"

    occasion_lines = "\n".join(
        f"  {occ_type:<38} {count:,}"
        for occ_type, count in list(occ["by_occasion_type"].items())[:10]
    )

    avail_lines = "\n".join(
        f"  {status:<38} {count:,}"
        for status, count in report["availability"].items()
    )

    text = textwrap.dedent(f"""\
        Product Feed Report — {date_str}
        {'=' * 55}

        VOLUME
        ------
        Total rows:              {v['total_rows']:,}
        Unique family IDs:       {v['unique_parent_ids']:,}
        Unique IDs:              {v['unique_ids']:,}
        Total URLs:              {v['total_urls']:,}

        DEPARTMENTS (top 15 by URL count)
        ----------------------------------
          {'Department':<40} {'URLs':>5}  {'Mean':>8}  {'Median':>9}
        {dept_lines}

        DESCRIPTIONS  (assessed per unique ID: {d['unique_ids_assessed']:,})
        ------------
        Missing (no pre-content):    {d['missing_count']:,}
        Short  (< 150 words):        {d['short_count']:,}
        Good   (150–300 words):      {d['good_count']:,}
        Long   (> 300 words):        {d['long_count']:,}
        Cross-family duplicates:     {d['cross_parent_duplicate_count']:,}
        Families with mixed descs:   {d['family_description_issue_count']:,}

        Top cross-family duplicate descriptions:
        {dup_lines}

        Families with inconsistent descriptions (top 10):
        {family_issue_lines}

        PRICING
        -------
        Minimum price:           £{p['min']}
        Maximum price:           £{p['max']}
        Mean price:              £{p['mean']}
        Median price:            £{p['median']}
        Products with no price:  {p['no_price_count']:,}
        Products with sale price:{p['with_sale_price_count']:,}

        AVAILABILITY
        ------------
        {avail_lines}

        OCCASION / PERSONALISATION
        --------------------------
        Occasion products total: {occ['is_occasion_count']:,}

        By occasion type:
        {occasion_lines}
    """)

    out_path = REPORTS_DIR / f"report_{date_str}.txt"
    out_path.write_text(text, encoding="utf-8")
    print(f"Text report saved to {out_path}")
    return out_path


# ── Excel helpers ──────────────────────────────────────────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SECTION_FILL = PatternFill("solid", fgColor="D6E4F0")
_SECTION_FONT = Font(bold=True)


def _hdr(ws, row, col, text, width=None):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = _HEADER_FONT
    cell.fill = _HEADER_FILL
    cell.alignment = Alignment(horizontal="left")
    if width:
        ws.column_dimensions[cell.column_letter].width = width


def _section(ws, row, text, ncols=2):
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    if ncols > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    return row + 1


def _kv(ws, row, label, value):
    ws.cell(row=row, column=1, value=label)
    ws.cell(row=row, column=2, value=value).alignment = Alignment(horizontal="right")
    return row + 1


# ── Main Excel export ──────────────────────────────────────────────────────────

def export_excel(report: dict, comparison: dict | None = None) -> Path:
    date_str = report["date"]
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    _hdr(ws, 1, 1, f"Product Feed Report — {date_str}", width=40)
    _hdr(ws, 1, 2, "Value", width=18)

    r = 2
    v = report["volume"]
    r = _section(ws, r, "Volume")
    r = _kv(ws, r, "Total rows", v["total_rows"])
    r = _kv(ws, r, "Unique family IDs", v["unique_parent_ids"])
    r = _kv(ws, r, "Unique IDs", v["unique_ids"])
    r = _kv(ws, r, "Total URLs", v["total_urls"])
    r += 1

    d = report["descriptions"]
    r = _section(ws, r, f"Descriptions  (unique IDs assessed: {d['unique_ids_assessed']:,})")
    r = _kv(ws, r, "Missing (no pre-content)", d["missing_count"])
    r = _kv(ws, r, "Short (< 150 words)", d["short_count"])
    r = _kv(ws, r, "Good (150–300 words)", d["good_count"])
    r = _kv(ws, r, "Long (> 300 words)", d["long_count"])
    r = _kv(ws, r, "Cross-family duplicates", d["cross_parent_duplicate_count"])
    r = _kv(ws, r, "Families with mixed descriptions", d["family_description_issue_count"])
    r += 1

    p = report["pricing"]
    r = _section(ws, r, "Pricing (overall)")
    r = _kv(ws, r, "Minimum price", p["min"])
    r = _kv(ws, r, "Maximum price", p["max"])
    r = _kv(ws, r, "Mean price", p["mean"])
    r = _kv(ws, r, "Median price", p["median"])
    r = _kv(ws, r, "No price", p["no_price_count"])
    r = _kv(ws, r, "With sale price", p["with_sale_price_count"])
    r += 1

    r = _section(ws, r, "Availability")
    for status, count in report["availability"].items():
        r = _kv(ws, r, status.title(), count)
    r += 1

    occ = report["occasion_products"]
    r = _section(ws, r, "Occasion / Personalisation")
    r = _kv(ws, r, "Total occasion products", occ["is_occasion_count"])

    # ── Sheet 2: Departments ──────────────────────────────────────────────────
    ws2 = wb.create_sheet("Departments")
    _hdr(ws2, 1, 1, "Department", width=44)
    _hdr(ws2, 1, 2, "URL Count", width=12)
    _hdr(ws2, 1, 3, "Mean Price (£)", width=16)
    _hdr(ws2, 1, 4, "Median Price (£)", width=16)
    for i, (dept, data) in enumerate(report["departments"].items(), start=2):
        ws2.cell(row=i, column=1, value=dept)
        ws2.cell(row=i, column=2, value=data["url_count"]).alignment = Alignment(horizontal="right")
        ws2.cell(row=i, column=3, value=data["price_mean"]).alignment = Alignment(horizontal="right")
        ws2.cell(row=i, column=4, value=data["price_median"]).alignment = Alignment(horizontal="right")

    # ── Sheet 3: ID Quality ───────────────────────────────────────────────────
    ws3 = wb.create_sheet("ID Quality")
    _hdr(ws3, 1, 1, "Family ID", width=22)
    _hdr(ws3, 1, 2, "ID", width=22)
    _hdr(ws3, 1, 3, "Quality", width=10)
    _hdr(ws3, 1, 4, "Word Count", width=12)
    _hdr(ws3, 1, 5, "Pre-content snippet", width=80)
    _hdr(ws3, 1, 6, "Example URL", width=60)
    for i, item in enumerate(report.get("id_quality_details", []), start=2):
        q = item["quality"]
        fill = PatternFill("solid", fgColor=_QUALITY_COLOURS.get(q, "FFFFFF"))
        for col, val in enumerate(
            [item["family_id"], item["id"], q, item["word_count"], item["pre_text_snippet"], item.get("example_url", "")],
            start=1,
        ):
            cell = ws3.cell(row=i, column=col, value=val)
            cell.fill = fill
            if col in (3, 4):
                cell.alignment = Alignment(horizontal="center")

    # ── Sheet 4: Family Description Issues ───────────────────────────────────
    ws4 = wb.create_sheet("Family Issues")
    _hdr(ws4, 1, 1, "Family ID", width=22)
    _hdr(ws4, 1, 2, "Distinct Descriptions", width=22)
    _hdr(ws4, 1, 3, "Example URL", width=60)
    _hdr(ws4, 1, 4, "Sample 1", width=60)
    _hdr(ws4, 1, 5, "Sample 2", width=60)
    _hdr(ws4, 1, 6, "Sample 3", width=60)
    issues = d.get("family_description_issues", [])
    for i, issue in enumerate(issues, start=2):
        ws4.cell(row=i, column=1, value=issue["family_id"])
        ws4.cell(row=i, column=2, value=issue["distinct_description_count"]).alignment = Alignment(horizontal="center")
        ws4.cell(row=i, column=3, value=issue.get("example_url", ""))
        for col, sample in enumerate(issue["pre_text_samples"][:3], start=4):
            ws4.cell(row=i, column=col, value=sample)

    # ── Sheet 5: Duplicate Descriptions ──────────────────────────────────────
    ws5 = wb.create_sheet("Duplicate Descriptions")
    _hdr(ws5, 1, 1, "Family ID Count", width=16)
    _hdr(ws5, 1, 2, "Pre-content (before 'Hamper contents include')", width=100)
    _hdr(ws5, 1, 3, "Family IDs (sample)", width=50)
    _hdr(ws5, 1, 4, "Example URL", width=60)
    for i, ex in enumerate(d["cross_parent_duplicate_examples"], start=2):
        ws5.cell(row=i, column=1, value=ex["family_id_count"]).alignment = Alignment(horizontal="right")
        ws5.cell(row=i, column=2, value=ex["pre_text"] or "(no pre-content)")
        ws5.cell(row=i, column=3, value=", ".join(ex["family_ids"][:5]))
        ws5.cell(row=i, column=4, value=ex.get("example_url", ""))

    # ── Sheet 6: Occasion Types ───────────────────────────────────────────────
    ws6 = wb.create_sheet("Occasion Types")
    _hdr(ws6, 1, 1, "Occasion Type", width=38)
    _hdr(ws6, 1, 2, "Count", width=14)
    for i, (occ_type, count) in enumerate(occ["by_occasion_type"].items(), start=2):
        ws6.cell(row=i, column=1, value=occ_type)
        ws6.cell(row=i, column=2, value=count).alignment = Alignment(horizontal="right")

    # ── Sheet 7: Day-on-day comparison ───────────────────────────────────────
    if comparison:
        ws7 = wb.create_sheet("Day-on-Day")
        _hdr(ws7, 1, 1, "Metric", width=40)
        _hdr(ws7, 1, 2, comparison["yesterday"], width=14)
        _hdr(ws7, 1, 3, comparison["today"], width=14)
        _hdr(ws7, 1, 4, "Change", width=18)

        u = comparison["urls"]
        r2 = 2
        r2 = _section(ws7, r2, "URLs & Rows", ncols=4)
        for label, y_val, t_val, chg in [
            ("Total URLs", u["yesterday"], u["today"], f"{u['change']:+,} ({u['change_pct']:+.1f}%)"),
            ("Total Rows", comparison["rows"]["yesterday"], comparison["rows"]["today"], f"{comparison['rows']['change']:+,}"),
        ]:
            ws7.cell(row=r2, column=1, value=label)
            ws7.cell(row=r2, column=2, value=y_val).alignment = Alignment(horizontal="right")
            ws7.cell(row=r2, column=3, value=t_val).alignment = Alignment(horizontal="right")
            ws7.cell(row=r2, column=4, value=chg).alignment = Alignment(horizontal="right")
            r2 += 1

        if u["dramatic_change"]:
            cell = ws7.cell(row=r2, column=1, value="⚠ DRAMATIC URL CHANGE DETECTED")
            cell.font = Font(bold=True, color="FF0000")
            r2 += 1

        r2 += 1
        ids = comparison["ids"]
        r2 = _section(ws7, r2, "Product IDs", ncols=4)
        ws7.cell(row=r2, column=1, value="New IDs")
        ws7.cell(row=r2, column=2, value=ids["new_count"]).alignment = Alignment(horizontal="right")
        r2 += 1
        ws7.cell(row=r2, column=1, value="Removed IDs")
        ws7.cell(row=r2, column=2, value=ids["removed_count"]).alignment = Alignment(horizontal="right")
        r2 += 2

        new_fam = comparison.get("new_family_ids", {})
        if new_fam.get("count", 0) > 0:
            r2 = _section(ws7, r2, f"New Family IDs ({new_fam['count']})", ncols=4)
            _hdr(ws7, r2, 1, "Family ID")
            _hdr(ws7, r2, 2, "Example Title")
            _hdr(ws7, r2, 3, "Example URL")
            r2 += 1
            for entry in new_fam.get("examples", []):
                ws7.cell(row=r2, column=1, value=entry["family_id"])
                ws7.cell(row=r2, column=2, value=entry.get("title", ""))
                ws7.cell(row=r2, column=3, value=entry.get("url", ""))
                r2 += 1
            r2 += 1

        if comparison.get("removed_family_ids", {}).get("count", 0) > 0:
            rem_fam = comparison["removed_family_ids"]
            r2 = _section(ws7, r2, f"Removed Family IDs ({rem_fam['count']})", ncols=4)
            for fid in rem_fam.get("ids", []):
                ws7.cell(row=r2, column=1, value=fid)
                r2 += 1
            r2 += 1

        if comparison.get("department_changes"):
            r2 = _section(ws7, r2, "Department Changes (URL count)", ncols=4)
            for dept, ch in list(comparison["department_changes"].items())[:20]:
                ws7.cell(row=r2, column=1, value=dept)
                ws7.cell(row=r2, column=2, value=ch["yesterday"]).alignment = Alignment(horizontal="right")
                ws7.cell(row=r2, column=3, value=ch["today"]).alignment = Alignment(horizontal="right")
                ws7.cell(row=r2, column=4, value=f"{ch['change']:+,}").alignment = Alignment(horizontal="right")
                r2 += 1

    out_path = REPORTS_DIR / f"report_{date_str}.xlsx"
    wb.save(out_path)
    print(f"Excel report saved to {out_path}")
    return out_path


# ── Comparison text / email report ────────────────────────────────────────────

def export_comparison_text(comp: dict) -> Path:
    u = comp["urls"]
    r = comp["rows"]
    ids = comp["ids"]
    nf = comp["new_family_ids"]
    rf = comp["removed_family_ids"]
    d = comp["descriptions"]
    p = comp["pricing"]

    alert = u["dramatic_change"]
    subject = (
        f"SUBJECT: RH Product Feed {comp['today']} vs {comp['yesterday']}"
        + (" [ALERT: DRAMATIC URL CHANGE]" if alert else "")
    )

    dept_lines = "\n".join(
        f"  {dept:<40} {ch['yesterday']:>5,} → {ch['today']:>5,}  ({ch['change']:+,})"
        for dept, ch in list(comp["department_changes"].items())[:15]
    ) or "  No department changes."

    new_examples = "\n".join(
        f"  {e['family_id']}: {e['title'][:60]}"
        + (f"\n    {e['url']}" if e.get("url") else "")
        for e in nf["examples"][:15]
    ) or "  None"

    removed_ids_sample = "\n".join(
        f"  {fid}" for fid in rf["ids"][:15]
    ) or "  None"

    text = textwrap.dedent(f"""\
        {subject}

        Product feed comparison: {comp['today']} vs {comp['yesterday']}
        {'=' * 60}

        VOLUME
        ------
        URLs:  {u['yesterday']:,} → {u['today']:,}  ({u['change']:+,} / {u['change_pct']:+.1f}%){"  *** DRAMATIC CHANGE ***" if alert else ""}
        Rows:  {r['yesterday']:,} → {r['today']:,}  ({r['change']:+,})

        PRODUCT CHANGES
        ---------------
        New product IDs:      {ids['new_count']:,}
        Removed product IDs:  {ids['removed_count']:,}
        New product families: {nf['count']:,}
        Removed families:     {rf['count']:,}

        New families (sample):
        {new_examples}

        Removed families (sample):
        {removed_ids_sample}

        DEPARTMENT CHANGES (URL count, top 15 by absolute change)
        ----------------------------------------------------------
        {'Department':<40} {'Yesterday':>9}   {'Today':>5}   Change
        {dept_lines}

        DESCRIPTION QUALITY
        -------------------
        Missing:               {d['missing_yesterday']:,} → {d['missing_today']:,}  ({d['missing_change']:+,})
        Cross-family dups:     {d['cross_parent_duplicates_yesterday']:,} → {d['cross_parent_duplicates_today']:,}
        Family inconsistencies:{d['family_issues_yesterday']:,} → {d['family_issues_today']:,}

        PRICING
        -------
        Mean:   £{p['mean_yesterday']} → £{p['mean_today']}
        Median: £{p['median_yesterday']} → £{p['median_today']}
    """)

    out_path = REPORTS_DIR / f"comparison_{comp['today']}_vs_{comp['yesterday']}.txt"
    out_path.write_text(text, encoding="utf-8")
    print(f"Comparison text report saved to {out_path}")
    return out_path
